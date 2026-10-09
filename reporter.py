"""Генератор отчетов для шефа: текстовый Telegram-формат и Excel (.xlsx)."""
import os
from typing import Dict, Any
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_telegram_text(data: Dict[str, Any]) -> str:
    """Генерирует аккуратный, представительный отчет для отправки шефу в Telegram."""
    period = data.get("period_title", "Отчетный период")
    total_msgs = data.get("total_messages", 0)
    client_msgs = data.get("client_messages_count", 0)
    mgr_msgs = data.get("manager_messages_count", 0)
    
    joins = data.get("joins_count", 0)
    leaves = data.get("leaves_count", 0)
    growth = data.get("net_growth", 0)
    growth_sign = "+" if growth > 0 else ""
    
    total_q = data.get("total_questions", 0)
    replies_count = data.get("total_replies", 0)
    avg_resp = data.get("avg_response_time", "0 сек")
    med_resp = data.get("median_response_time", "0 сек")
    pct_5m = data.get("pct_under_5m", 0)
    pct_15m = data.get("pct_under_15m", 0)
    pct_over_1h = data.get("pct_over_1h", 0)
    peak_hours = data.get("peak_hours_str", "—")
    
    text = (
        f"📊 <b>ОТЧЕТ ДЛЯ ШЕФА ПО ЧАТУ МАГАЗИНА</b>\n"
        f"🗓 <b>Период:</b> {period}\n"
        f"────────────────────────────\n\n"
        f"📈 <b>1. ОБЩАЯ АКТИВНОСТЬ:</b>\n"
        f"• Всего сообщений: <b>{total_msgs:,}</b>\n"
        f"  ├ от клиентов: <b>{client_msgs:,}</b>\n"
        f"  └ от менеджеров: <b>{mgr_msgs:,}</b>\n"
        f"• Новых участников: <b>+{joins}</b>\n"
        f"• Вышло из чата: <b>-{leaves}</b>\n"
        f"• Чистый прирост аудитории: <b>{growth_sign}{growth} чел.</b>\n"
        f"• Пиковые часы нагрузки: <i>{peak_hours}</i>\n\n"
        f"────────────────────────────\n"
        f"⚡ <b>2. СКОРОСТЬ ОТВЕТОВ ПОДДЕРЖКИ:</b>\n"
        f"• Зафиксировано ответов: <b>{replies_count}</b>\n"
        f"• Среднее время ответа: <b>{avg_resp}</b>\n"
        f"• Медианное время: <b>{med_resp}</b>\n"
        f"• Быстрых ответов (до 5 мин): <b>{pct_5m}%</b>\n"
        f"• В рамках нормы (до 15 мин): <b>{pct_15m}%</b>\n"
    )
    
    if pct_over_1h > 0:
        text += f"• ⚠️ Задержки дольше 1 часа: <b>{pct_over_1h}%</b>\n\n"
    else:
        text += f"• 🟢 Задержек свыше 1 часа: <b>0%</b> (отличная работа!)\n\n"
        
    text += (
        f"────────────────────────────\n"
        f"🎯 <b>3. ГЛАВНЫЕ «БОЛИ» И ВОПРОСЫ КЛИЕНТОВ ({total_q} вопр.):</b>\n"
    )
    
    cats = data.get("category_stats", [])
    if cats:
        for c in cats:
            text += (
                f"{c['emoji']} <b>{c['title']}</b>: {c['count']} ({c['pct']}%)\n"
                f"  <code>{c['progress_bar']}</code>\n"
            )
    else:
        text += "<i>За указанный период вопросов не зафиксировано.</i>\n"
        
    managers = data.get("managers_summary", [])
    if managers:
        text += (
            f"\n────────────────────────────\n"
            f"👥 <b>4. РЕЗУЛЬТАТЫ МЕНЕДЖЕРОВ:</b>\n"
        )
        for idx, m in enumerate(managers, 1):
            text += (
                f"{idx}. <b>{m['name']}</b>: "
                f"<b>{m['replies_count']}</b> отв. | среднее: <b>{m['avg_time_str']}</b>\n"
            )
            
    text += (
        f"\n────────────────────────────\n"
        f"💡 <i>Отчет сгенерирован автоматически в {datetime.now().strftime('%H:%M %d.%m.%Y')}.</i>\n"
        f"📎 Полная детализация с логами доступна в Excel-выгрузке."
    )
    
    return text


def generate_excel_file(data: Dict[str, Any], output_path: str) -> str:
    """Создает профессиональную Excel-таблицу со стилизацией и вкладками."""
    wb = openpyxl.Workbook()
    
    # Стили
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # ================= ЛИСТ 1: СВОДКА ДЛЯ ШЕФА =================
    ws_summary = wb.active
    ws_summary.title = "Сводка для шефа"
    ws_summary.views.sheetView[0].showGridLines = True
    
    ws_summary.cell(row=1, column=1, value="ОТЧЕТ ПО ЧАТУ МАГАЗИНА ДЛЯ ШЕФА").font = title_font
    ws_summary.cell(row=2, column=1, value=f"Период: {data.get('period_title')}").font = bold_font
    
    # KPI карточки
    kpis = [
        ("Всего сообщений", data.get("total_messages", 0)),
        ("Сообщений от клиентов", data.get("client_messages_count", 0)),
        ("Сообщений от менеджеров", data.get("manager_messages_count", 0)),
        ("Новых участников (прирост)", data.get("joins_count", 0)),
        ("Вышло участников", data.get("leaves_count", 0)),
        ("Всего обращений / вопросов", data.get("total_questions", 0)),
        ("Обработано ответов", data.get("total_replies", 0)),
        ("Средняя скорость ответа поддержки", data.get("avg_response_time", "0")),
        ("Медианное время ответа", data.get("median_response_time", "0")),
        ("Отвечено быстрее 5 минут", f"{data.get('pct_under_5m', 0)}%"),
        ("Отвечено быстрее 15 минут", f"{data.get('pct_under_15m', 0)}%"),
        ("Задержек ответа (> 1 часа)", f"{data.get('pct_over_1h', 0)}%"),
        ("Пиковые часы обращений", data.get("peak_hours_str", "—")),
    ]
    
    ws_summary.cell(row=4, column=1, value="Метрика").font = header_font
    ws_summary.cell(row=4, column=1).fill = header_fill
    ws_summary.cell(row=4, column=2, value="Значение").font = header_font
    ws_summary.cell(row=4, column=2).fill = header_fill
    
    for idx, (label, val) in enumerate(kpis, start=5):
        c1 = ws_summary.cell(row=idx, column=1, value=label)
        c2 = ws_summary.cell(row=idx, column=2, value=str(val))
        c1.font = regular_font
        c2.font = bold_font
        c1.border = thin_border
        c2.border = thin_border
        
    # Таблица категорий болей на листе сводки
    start_cat_row = len(kpis) + 7
    ws_summary.cell(row=start_cat_row, column=1, value="Категория вопроса (боль клиента)").font = header_font
    ws_summary.cell(row=start_cat_row, column=1).fill = header_fill
    ws_summary.cell(row=start_cat_row, column=2, value="Кол-во вопросов").font = header_font
    ws_summary.cell(row=start_cat_row, column=2).fill = header_fill
    ws_summary.cell(row=start_cat_row, column=3, value="Доля от всех (%)").font = header_font
    ws_summary.cell(row=start_cat_row, column=3).fill = header_fill
    
    for idx, c in enumerate(data.get("category_stats", []), start=start_cat_row + 1):
        c1 = ws_summary.cell(row=idx, column=1, value=c["title"])
        c2 = ws_summary.cell(row=idx, column=2, value=c["count"])
        c3 = ws_summary.cell(row=idx, column=3, value=f"{c['pct']}%")
        for cell in (c1, c2, c3):
            cell.font = regular_font
            cell.border = thin_border
        c2.alignment = center_align
        c3.alignment = center_align
        
    # ================= ЛИСТ 2: МЕНЕДЖЕРЫ =================
    ws_mgr = wb.create_sheet(title="Менеджеры")
    ws_mgr.views.sheetView[0].showGridLines = True
    
    mgr_headers = ["№", "Менеджер / Оператор", "Кол-во закрытых ответов", "Среднее время ответа", "Медианное время"]
    for col_idx, h in enumerate(mgr_headers, 1):
        cell = ws_mgr.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for row_idx, m in enumerate(data.get("managers_summary", []), start=2):
        ws_mgr.cell(row=row_idx, column=1, value=row_idx - 1).alignment = center_align
        ws_mgr.cell(row=row_idx, column=2, value=m["name"])
        ws_mgr.cell(row=row_idx, column=3, value=m["replies_count"]).alignment = center_align
        ws_mgr.cell(row=row_idx, column=4, value=m["avg_time_str"]).alignment = center_align
        ws_mgr.cell(row=row_idx, column=5, value=m["median_time_str"]).alignment = center_align
        for c in range(1, 6):
            cell = ws_mgr.cell(row=row_idx, column=c)
            cell.font = regular_font
            cell.border = thin_border

    # ================= ЛИСТ 3: ЛОГ ОТВЕТОВ =================
    ws_log = wb.create_sheet(title="Замеры времени ответов")
    ws_log.views.sheetView[0].showGridLines = True
    
    log_headers = ["ID", "Вопрос задан", "Ответ дан", "Клиент", "Менеджер", "Категория", "Время ответа (сек)", "Время ответа (формат)"]
    for col_idx, h in enumerate(log_headers, 1):
        cell = ws_log.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    for row_idx, r in enumerate(data.get("raw_responses", [])[:1000], start=2):
        from analyzer import format_duration
        dur_s = r.get("response_time_seconds", 0)
        ws_log.cell(row=row_idx, column=1, value=row_idx - 1).alignment = center_align
        ws_log.cell(row=row_idx, column=2, value=r.get("question_time", ""))
        ws_log.cell(row=row_idx, column=3, value=r.get("reply_time", ""))
        ws_log.cell(row=row_idx, column=4, value=r.get("client_name", ""))
        ws_log.cell(row=row_idx, column=5, value=r.get("manager_name", ""))
        ws_log.cell(row=row_idx, column=6, value=r.get("category", ""))
        ws_log.cell(row=row_idx, column=7, value=round(dur_s, 1)).alignment = center_align
        ws_log.cell(row=row_idx, column=8, value=format_duration(dur_s)).alignment = center_align
        for c in range(1, 9):
            cell = ws_log.cell(row=row_idx, column=c)
            cell.font = regular_font
            cell.border = thin_border

    # Авто-подгон ширины колонок
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = val.split("\n")[0]
                max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
